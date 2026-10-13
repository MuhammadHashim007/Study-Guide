import io
import json
import random
import re
import time
from dataclasses import dataclass
from datetime import datetime

import streamlit as st
from docx import Document
from openai import (
    APIConnectionError,
    APIStatusError,
    AuthenticationError,
    OpenAI,
    RateLimitError,
)
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

APP_TITLE = "OmniStudy AI - University Academic & Research Suite"
OPENROUTER_BASE_URL = "https://openrouter.ai/api/v1"
GOOGLE_BASE_URL = "https://generativelanguage.googleapis.com/v1beta/openai/"
FREE_ROUTER = "openrouter/free"
LETTERS = "ABCDEF"
WORDS_PER_MINUTE = 238

MODEL_CATALOG = {
    "Frontier Academic & Reasoning": [
        {
            "label": "Anthropic: Claude 3.5 Sonnet",
            "id": "anthropic/claude-3.5-sonnet",
            "provider": "openrouter",
            "tier": "Paid",
            "note": "Supreme academic prose, nuanced paper synthesis.",
        },
        {
            "label": "OpenAI: GPT-4o",
            "id": "openai/gpt-4o",
            "provider": "openrouter",
            "tier": "Paid",
            "note": "Multimodal flagship, precise structured schemas.",
        },
        {
            "label": "DeepSeek: R1",
            "id": "deepseek/deepseek-r1",
            "provider": "openrouter",
            "tier": "Paid",
            "note": "Chain-of-thought STEM and mathematical logic.",
        },
        {
            "label": "DeepSeek: R1 (free)",
            "id": "deepseek/deepseek-r1:free",
            "provider": "openrouter",
            "tier": "Free",
            "note": "Same reasoning model on free capacity. Slow, rate-limited.",
        },
        {
            "label": "Google: Gemini Pro 1.5",
            "id": "google/gemini-pro-1.5",
            "provider": "openrouter",
            "tier": "Paid",
            "note": "Very large context window for textbooks and slide sets.",
        },
        {
            "label": "Google: Gemini 2.0 Flash Exp (free)",
            "id": "google/gemini-2.0-flash-exp:free",
            "provider": "openrouter",
            "tier": "Free",
            "note": "Large context on free capacity. Experimental endpoints get retired often.",
        },
    ],
    "Open-Weights Powerhouses": [
        {
            "label": "Meta: Llama 3.3 70B Instruct (free)",
            "id": "meta-llama/llama-3.3-70b-instruct:free",
            "provider": "openrouter",
            "tier": "Free",
            "note": "Strong general-purpose open model.",
        },
        {
            "label": "Meta: Llama 3.3 70B Instruct",
            "id": "meta-llama/llama-3.3-70b-instruct",
            "provider": "openrouter",
            "tier": "Paid",
            "note": "Same model without free-tier rate limits.",
        },
        {
            "label": "Qwen: Qwen 2.5 72B Instruct (free)",
            "id": "qwen/qwen-2.5-72b-instruct:free",
            "provider": "openrouter",
            "tier": "Free",
            "note": "Good at structured output and STEM.",
        },
        {
            "label": "Qwen: Qwen 2.5 72B Instruct",
            "id": "qwen/qwen-2.5-72b-instruct",
            "provider": "openrouter",
            "tier": "Paid",
            "note": "Same model without free-tier rate limits.",
        },
        {
            "label": "Mistral: Large",
            "id": "mistralai/mistral-large",
            "provider": "openrouter",
            "tier": "Paid",
            "note": "Reliable multilingual academic writing.",
        },
        {
            "label": "Mistral: Nemo (free)",
            "id": "mistralai/mistral-nemo:free",
            "provider": "openrouter",
            "tier": "Free",
            "note": "Small and fast. Weakest on long or technical material.",
        },
    ],
    "Zero-Cost Free Auto-Router": [
        {
            "label": "OpenRouter Free Auto-Router",
            "id": FREE_ROUTER,
            "provider": "openrouter",
            "tier": "Free",
            "note": "Routes to whichever free model currently has capacity. Most reliable free option, quality varies per request.",
        },
    ],
    "Direct Google AI Studio (own key)": [
        {
            "label": "Google: Gemini 2.5 Flash (direct)",
            "id": "gemini-2.5-flash",
            "provider": "google",
            "tier": "Free tier",
            "note": "Uses your Google AI Studio key directly. Large context, generous free quota.",
        },
        {
            "label": "Google: Gemini 2.5 Pro (direct)",
            "id": "gemini-2.5-pro",
            "provider": "google",
            "tier": "Limited",
            "note": "Strongest Gemini. Free-tier access is tight or absent, check your quota.",
        },
    ],
}

SYSTEM_BASE = (
    "You are an expert university tutor and academic writer. Work only from the "
    "SOURCE MATERIAL supplied by the student. When you add background knowledge that "
    "is not in the source, label it clearly as (outside source). Never invent "
    "citations, statistics, quotations, or bibliographic details. If something is "
    "missing from the source, say so explicitly. Write mathematical expressions in "
    "LaTeX using $...$ for inline and $$...$$ for display math. Use clean Markdown."
)

SUMMARY_FORMATS = {
    "High-Yield Executive Summary": (
        "Produce a high-yield executive summary of at most about 500 words: "
        "(1) a three-sentence TL;DR, (2) the core concepts with one-line definitions, "
        "(3) key formulas, laws or frameworks, (4) the five most examinable points, "
        "(5) open questions or gaps in the source."
    ),
    "Comprehensive Exam Revision Guide": (
        "Produce a comprehensive exam revision guide organised by topic. For each topic "
        "give: definitions, mechanisms or derivations, worked examples drawn from the "
        "source, common mistakes, and a mnemonic where useful. End with a one-page quick "
        "review sheet and 8 self-test questions with short model answers."
    ),
    "First-Principles Conceptual Explainer": (
        "Explain the material from first principles. Start with the simplest "
        "assumptions, build up in layers, use at least two concrete analogies, name the "
        "most common misconceptions and why they are wrong, and finish with "
        "'check your understanding' questions that test reasoning rather than recall."
    ),
}

QUIZ_DIFFICULTY = {
    "Undergraduate Exam": "standard undergraduate exam level: application of concepts, not just recall",
    "Graduate/Complex": "graduate level: multi-step reasoning, edge cases, subtle distractors and synthesis across topics",
    "Foundational": "foundational level: core definitions and basic concept recognition",
}

PROJECT_TIMELINES = [
    "1 week",
    "2 weeks",
    "4 weeks",
    "8 weeks",
    "12 weeks",
    "Full semester (15 weeks)",
]

NO_SHUFFLE_PATTERN = re.compile(
    r"\b(all|none|both|neither)\s+of\s+the\s+(above|these)\b"
    r"|\bboth\s+[A-D]\s+and\s+[A-D]\b"
    r"|\b[A-D]\s+and\s+[A-D]\s+(only|are)\b",
    re.IGNORECASE,
)


class LLMError(Exception):
    pass


@dataclass
class ModelConfig:
    provider: str
    model: str
    label: str
    api_key: str
    fallback: bool


@dataclass
class Source:
    text: str
    words: int
    chars: int
    truncated: bool


def get_secret(name):
    try:
        value = st.secrets[name]
    except Exception:
        return ""
    return str(value).strip()


def strip_reasoning(text):
    return re.sub(r"<think>.*?(</think>|$)", "", text or "", flags=re.DOTALL | re.IGNORECASE)


def slugify(text):
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    return slug or "omnistudy"


def shape_lines(shape):
    lines = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            lines.extend(shape_lines(child))
        return lines
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs).strip()
            if text:
                indent = "  " * min(paragraph.level, 4)
                lines.append(f"{indent}- {text}")
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append("  | " + " | ".join(cells) + " |")
    return lines


def extract_pdf(data):
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception:
            unlocked = 0
        if not unlocked:
            return "", "PDF is password-protected and could not be opened."
    parts = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            parts.append(f"[Page {number}]\n{text.strip()}")
    combined = "\n\n".join(parts)
    if not combined.strip():
        return "", "No extractable text found. This is probably a scanned PDF; run OCR first or paste the text."
    return combined, ""


def extract_docx(data):
    document = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    combined = "\n".join(parts)
    if not combined.strip():
        return "", "The document contains no extractable text."
    return combined, ""


def extract_pptx(data):
    presentation = Presentation(io.BytesIO(data))
    parts = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines = []
        for shape in slide.shapes:
            lines.extend(shape_lines(shape))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        block = [f"[Slide {number}]"] + lines
        if notes:
            block.append(f"Speaker notes: {notes}")
        if len(block) > 1:
            parts.append("\n".join(block))
    combined = "\n\n".join(parts)
    if not combined.strip():
        return "", "No text found on any slide (image-only slides cannot be read)."
    return combined, ""


def extract_txt(data):
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except (UnicodeDecodeError, UnicodeError):
            text = ""
    if not text.strip():
        return "", "The text file is empty or could not be decoded."
    return text, ""


@st.cache_data(show_spinner=False, max_entries=32)
def extract_text(name, data):
    extension = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    handlers = {
        "pdf": extract_pdf,
        "docx": extract_docx,
        "pptx": extract_pptx,
        "txt": extract_txt,
    }
    handler = handlers.get(extension)
    if handler is None:
        return "", f"Unsupported file type '.{extension}'. Use PDF, DOCX, PPTX or TXT."
    try:
        return handler(data)
    except Exception as exc:
        return "", f"Could not read this file ({type(exc).__name__}). It may be corrupted."


def prepare_source(text, max_chars):
    words = len(text.split())
    chars = len(text)
    if chars <= max_chars:
        return Source(text, words, chars, False)
    cut = text[:max_chars]
    boundary = max(cut.rfind("\n"), cut.rfind(" "))
    if boundary > max_chars * 0.9:
        cut = cut[:boundary]
    return Source(cut, words, chars, True)


def build_client(cfg):
    if cfg.provider == "google":
        return OpenAI(api_key=cfg.api_key, base_url=GOOGLE_BASE_URL, timeout=180, max_retries=1)
    return OpenAI(
        api_key=cfg.api_key,
        base_url=OPENROUTER_BASE_URL,
        timeout=180,
        max_retries=1,
        default_headers={"HTTP-Referer": "https://streamlit.io", "X-Title": "OmniStudy AI"},
    )


def error_detail(exc):
    message = getattr(exc, "message", None) or str(exc)
    return message[:300]


def message_text(response):
    choices = getattr(response, "choices", None)
    if not choices:
        return ""
    message = getattr(choices[0], "message", None)
    content = getattr(message, "content", None) or ""
    return strip_reasoning(content).strip()


def call_llm(cfg, system, user, temperature=0.3, max_tokens=4096):
    if not cfg.api_key:
        raise LLMError(
            "No API key found. Add OPENROUTER_API_KEY (or GOOGLE_API_KEY for direct Gemini) "
            "to Streamlit Secrets, or paste a key in the sidebar."
        )
    client = build_client(cfg)
    candidates = [cfg.model]
    if cfg.provider == "openrouter" and cfg.fallback and cfg.model != FREE_ROUTER:
        candidates.append(FREE_ROUTER)
    last_error = None
    for model in candidates:
        for attempt in range(2):
            try:
                response = client.chat.completions.create(
                    model=model,
                    messages=[
                        {"role": "system", "content": system},
                        {"role": "user", "content": user},
                    ],
                    temperature=temperature,
                    max_tokens=max_tokens,
                )
                text = message_text(response)
                if text:
                    return text, model
                last_error = LLMError(f"Model '{model}' returned an empty response.")
                break
            except AuthenticationError:
                raise LLMError("HTTP 401: the API key was rejected. Check the key and that it belongs to the selected provider.")
            except RateLimitError as exc:
                last_error = LLMError(
                    f"HTTP 429: rate limit or quota reached for '{model}'. Free models are throttled "
                    f"heavily; wait a minute, pick another model, or use the Auto-Router. ({error_detail(exc)})"
                )
                if attempt == 0:
                    time.sleep(4)
                    continue
                break
            except APIConnectionError:
                raise LLMError("Network error: could not reach the model provider. Check your connection and retry.")
            except APIStatusError as exc:
                status = exc.status_code
                if status == 402:
                    last_error = LLMError(
                        f"HTTP 402: insufficient credits for '{model}'. Add credits or choose a free model."
                    )
                    break
                if status == 404:
                    last_error = LLMError(
                        f"HTTP 404: '{model}' is unavailable. It may have been retired, or your OpenRouter "
                        f"privacy settings block free endpoints. ({error_detail(exc)})"
                    )
                    break
                last_error = LLMError(f"HTTP {status} from '{model}': {error_detail(exc)}")
                if status in (500, 502, 503, 504) and attempt == 0:
                    time.sleep(2)
                    continue
                break
    raise last_error or LLMError("The request failed for an unknown reason.")


def find_objects(raw, required_keys):
    text = strip_reasoning(raw)
    text = re.sub(r"```(?:json|JSON)?", "", text)
    decoder = json.JSONDecoder()
    found = []
    position = 0
    while True:
        start = text.find("{", position)
        if start == -1:
            break
        try:
            obj, length = decoder.raw_decode(text[start:])
        except json.JSONDecodeError:
            position = start + 1
            continue
        if isinstance(obj, dict) and required_keys.issubset(obj.keys()):
            found.append(obj)
            position = start + length
        else:
            position = start + 1
    return found


def clean_option(text):
    return re.sub(r"^\s*\(?[A-Fa-f][\.\)\:]\s+", "", str(text)).strip()


def resolve_correct(value, texts, keys):
    count = len(texts)
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if 0 <= value < count else None
    if isinstance(value, str):
        candidate = value.strip()
        for index, option in enumerate(texts):
            if option.strip().lower() == candidate.lower():
                return index
        match = re.match(r"^\(?([A-Fa-f])\)?(?:[\.\):\s]|$)", candidate)
        if match:
            index = ord(match.group(1).upper()) - 65
            if 0 <= index < count:
                return index
        if keys and candidate in keys:
            return keys.index(candidate)
    return None


def normalize_rationales(raw, count, keys, correct):
    output = [""] * count
    if isinstance(raw, list):
        for index, item in enumerate(raw[:count]):
            output[index] = str(item).strip()
    elif isinstance(raw, dict):
        for index in range(count):
            for key in ((keys[index] if keys else None), LETTERS[index]):
                if key is not None and key in raw:
                    output[index] = str(raw[key]).strip()
                    break
    elif isinstance(raw, str) and raw.strip():
        output[correct] = raw.strip()
    return output


def normalize_question(obj):
    question = str(obj.get("question", "")).strip()
    options = obj.get("options")
    keys = []
    if isinstance(options, dict):
        keys = list(options.keys())
        texts = [clean_option(options[k]) for k in keys]
    elif isinstance(options, list):
        texts = [clean_option(item) for item in options]
    else:
        return None
    texts = [t for t in texts if t]
    if not question or not 2 <= len(texts) <= len(LETTERS):
        return None
    correct_raw = obj.get("correct", obj.get("answer", obj.get("correct_answer")))
    correct = resolve_correct(correct_raw, texts, keys)
    if correct is None:
        return None
    rationale_raw = obj.get("rationales", obj.get("explanations", obj.get("explanation")))
    rationales = normalize_rationales(rationale_raw, len(texts), keys, correct)
    if not any(NO_SHUFFLE_PATTERN.search(t) for t in texts):
        order = list(range(len(texts)))
        random.shuffle(order)
        texts = [texts[i] for i in order]
        rationales = [rationales[i] for i in order]
        correct = order.index(correct)
    return {"question": question, "options": texts, "correct": correct, "rationales": rationales}


def normalize_card(obj):
    front = str(obj.get("front", "")).strip()
    back = str(obj.get("back", "")).strip()
    if not front or not back:
        return None
    return {"front": front, "back": back, "topic": str(obj.get("topic", "")).strip()}


def option_label(index, text):
    return f"{LETTERS[index]}. {text}"


def with_source(source_text, task):
    return f'SOURCE MATERIAL (between triple quotes):\n"""\n{source_text}\n"""\n\nTASK:\n{task}'


def build_markdown(title, body, model):
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    return f"# {title}\n\n_Generated by OmniStudy AI with `{model}` on {stamp}_\n\n---\n\n{body}\n"


def table_cell(text):
    return text.replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


def flashcards_to_table(cards):
    lines = ["| # | Front | Back |", "|---|-------|------|"]
    for number, card in enumerate(cards, start=1):
        lines.append(f"| {number} | {table_cell(card['front'])} | {table_cell(card['back'])} |")
    return "\n".join(lines)


def flashcards_to_tsv(cards):
    rows = []
    for card in cards:
        front = card["front"].replace("\t", " ").replace("\n", "<br>")
        back = card["back"].replace("\t", " ").replace("\n", "<br>")
        rows.append(f"{front}\t{back}")
    return "\n".join(rows)


def quiz_to_markdown(quiz, answers=None):
    lines = []
    for number, q in enumerate(quiz, start=1):
        lines.append(f"### Q{number}. {q['question']}\n")
        for index, text in enumerate(q["options"]):
            lines.append(f"- {option_label(index, text)}")
        lines.append("")
    lines.append("## Answer Key and Rationales\n")
    for number, q in enumerate(quiz, start=1):
        chosen = ""
        if answers is not None:
            picked = answers[number - 1]
            chosen = f" (your answer: {picked if picked else 'none'})"
        lines.append(f"**Q{number}: {LETTERS[q['correct']]}**{chosen}\n")
        for index, text in enumerate(q["options"]):
            marker = "CORRECT" if index == q["correct"] else "Incorrect"
            rationale = q["rationales"][index] or "No rationale provided."
            lines.append(f"- {LETTERS[index]} ({marker}): {rationale}")
        lines.append("")
    return "\n".join(lines)


def init_state():
    defaults = {
        "summary": None,
        "notes": None,
        "paper": None,
        "project": None,
        "cards": None,
        "cards_meta": None,
        "fc_queue": None,
        "fc_reveal": False,
        "fc_good": 0,
        "fc_again": 0,
        "quiz": None,
        "quiz_meta": None,
        "quiz_submitted": False,
        "quiz_answers": None,
    }
    for key, value in defaults.items():
        st.session_state.setdefault(key, value)


def inject_css():
    st.markdown(
        """
        <style>
        .block-container {padding-top: 2rem;}
        div[data-testid="stMetric"] {
            background: rgba(128,128,128,0.08);
            border: 1px solid rgba(128,128,128,0.25);
            padding: 12px 16px;
            border-radius: 10px;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


def render_sidebar():
    st.sidebar.title("🎓 OmniStudy AI")
    st.sidebar.caption("University Academic & Research Suite")
    st.sidebar.subheader("Model")
    category = st.sidebar.selectbox("Category", list(MODEL_CATALOG.keys()))
    entry = st.sidebar.selectbox(
        "Model",
        MODEL_CATALOG[category],
        format_func=lambda e: f"{e['label']} [{e['tier']}]",
    )
    st.sidebar.caption(entry["note"])
    custom = st.sidebar.text_input(
        "Custom model ID (optional)",
        placeholder="provider/model-name",
        help="Overrides the model above. Useful when a listed model has been retired. Sent to the same provider as the selected category.",
    )
    provider = entry["provider"]
    model = custom.strip() or entry["id"]
    label = custom.strip() or entry["label"]

    st.sidebar.subheader("API keys")
    openrouter_input = st.sidebar.text_input("OpenRouter API Key", type="password")
    google_input = st.sidebar.text_input("Google AI Studio API Key (optional)", type="password")
    openrouter_key = openrouter_input.strip() or get_secret("OPENROUTER_API_KEY")
    google_key = google_input.strip() or get_secret("GOOGLE_API_KEY")
    api_key = google_key if provider == "google" else openrouter_key
    needed = "Google AI Studio" if provider == "google" else "OpenRouter"
    if api_key:
        st.sidebar.success(f"{needed} key loaded")
    else:
        st.sidebar.warning(f"No {needed} key found")

    st.sidebar.subheader("Advanced")
    max_chars = st.sidebar.slider(
        "Max characters sent to model",
        min_value=10000,
        max_value=400000,
        value=60000,
        step=10000,
        help="About 4 characters per token. Free models often cap context well below 128k tokens.",
    )
    fallback = True
    if provider == "openrouter" and model != FREE_ROUTER:
        fallback = st.sidebar.checkbox(
            "Fall back to the free Auto-Router if the model fails",
            value=True,
            help="Triggers on 402, 404, 429 and server errors. The result is labelled with the model that actually answered.",
        )
    st.sidebar.caption(
        "Free OpenRouter endpoints may log prompts. Do not upload confidential or unpublished research to them."
    )
    return ModelConfig(provider, model, label, api_key, fallback), max_chars


def render_ingestion(max_chars):
    st.subheader("1. Ingest your material")
    files = st.file_uploader(
        "Upload lecture slides, papers, handouts or notes",
        type=["pdf", "docx", "pptx", "txt"],
        accept_multiple_files=True,
    )
    with st.expander("✍️ Or paste lecture notes, a syllabus, or a draft paper", expanded=False):
        pasted = st.text_area(
            "Pasted text",
            height=320,
            label_visibility="collapsed",
            placeholder="Paste text here. It is combined with any uploaded files.",
        )
    parts = []
    report = []
    for uploaded in files or []:
        text, warning = extract_text(uploaded.name, uploaded.getvalue())
        if warning:
            st.warning(f"{uploaded.name}: {warning}")
        if text.strip():
            parts.append(f"=== FILE: {uploaded.name} ===\n{text.strip()}")
            report.append({"File": uploaded.name, "Words": len(text.split()), "Characters": len(text)})
    if pasted.strip():
        parts.append(f"=== PASTED TEXT ===\n{pasted.strip()}")
        report.append({"File": "Pasted text", "Words": len(pasted.split()), "Characters": len(pasted)})
    combined = "\n\n".join(parts)
    source = prepare_source(combined, max_chars)
    if report:
        st.dataframe(report, hide_index=True)
        st.caption(
            f"Total: {source.chars:,} characters, {source.words:,} words, "
            f"about {source.chars // 4:,} tokens."
        )
        if source.truncated:
            st.warning(
                f"Material exceeds the {max_chars:,}-character limit. Only the first part is sent to the model. "
                "Raise the limit in the sidebar if your model supports a larger context, or split the material."
            )
    return source


def render_metrics(cfg, source):
    reading = source.words / WORDS_PER_MINUTE
    revision = reading * 2.5
    col1, col2, col3 = st.columns(3)
    col1.metric("Active Model", cfg.label)
    col2.metric("Ingested Words", f"{source.words:,}")
    col3.metric("Est. Reading / Revision", f"{reading:.0f} / {revision:.0f} min")


def ensure_ready(source):
    if not source.text.strip():
        st.warning("Upload a file or paste some text first.")
        return False
    return True


def generate_text(cfg, state_key, system, user, max_tokens=4096, temperature=0.3):
    try:
        with st.spinner(f"Generating with {cfg.label}..."):
            text, used = call_llm(cfg, system, user, temperature=temperature, max_tokens=max_tokens)
    except LLMError as exc:
        st.error(str(exc))
        return
    st.session_state[state_key] = {"text": text, "model": used, "requested": cfg.model}


def show_text_result(state_key, title, filename_stub, button_key):
    result = st.session_state.get(state_key)
    if not result:
        return
    if result["model"] != result["requested"]:
        st.info(f"The requested model failed, so `{result['model']}` answered instead.")
    st.markdown(result["text"])
    st.download_button(
        "⬇️ Download as Markdown",
        data=build_markdown(title, result["text"], result["model"]),
        file_name=f"{filename_stub}.md",
        mime="text/markdown",
        key=button_key,
    )


def render_summary_tab(cfg, source):
    st.markdown("Condense your material into the format that fits your study goal.")
    format_name = st.selectbox("Summary format", list(SUMMARY_FORMATS.keys()), key="summary_format")
    focus = st.text_input("Focus on (optional)", placeholder="e.g. chapters 3-4, thermodynamics, exam on Friday", key="summary_focus")
    if st.button("Generate summary", type="primary", key="summary_go"):
        if ensure_ready(source):
            task = SUMMARY_FORMATS[format_name]
            if focus.strip():
                task += f"\n\nPrioritise this focus area: {focus.strip()}"
            generate_text(cfg, "summary", SYSTEM_BASE, with_source(source.text, task), max_tokens=4096)
    show_text_result("summary", f"{format_name}", f"summary-{slugify(format_name)}", "summary_dl")


def start_drill(total):
    st.session_state["fc_queue"] = list(range(total))
    st.session_state["fc_reveal"] = False
    st.session_state["fc_good"] = 0
    st.session_state["fc_again"] = 0


def reveal_card():
    st.session_state["fc_reveal"] = True


def grade_card(good):
    queue = st.session_state["fc_queue"]
    if not queue:
        return
    index = queue.pop(0)
    if good:
        st.session_state["fc_good"] += 1
    else:
        st.session_state["fc_again"] += 1
        queue.insert(min(3, len(queue)), index)
    st.session_state["fc_reveal"] = False


def render_drill(cards):
    st.markdown("#### 🔁 Spaced-repetition drill")
    st.caption(
        "Cards you miss come back after a few other cards in the same session. "
        "For long-term retention, repeat the full deck after 1, 3, 7, 14 and 30 days (Leitner schedule)."
    )
    st.button("Start / restart drill", key="fc_start", on_click=start_drill, args=(len(cards),))
    queue = st.session_state["fc_queue"]
    if queue is None:
        return
    if not queue:
        st.success(
            f"Drill complete. Recalled: {st.session_state['fc_good']} · Missed attempts: {st.session_state['fc_again']}"
        )
        return
    card = cards[queue[0]]
    st.progress(
        st.session_state["fc_good"] / len(cards),
        text=f"{st.session_state['fc_good']}/{len(cards)} mastered · {len(queue)} cards in queue",
    )
    st.markdown(f"**Front:** {card['front']}")
    if not st.session_state["fc_reveal"]:
        st.button("Reveal answer", key="fc_reveal_btn", on_click=reveal_card)
    else:
        st.markdown(f"**Back:** {card['back']}")
        col1, col2 = st.columns(2)
        col1.button("❌ Again", key="fc_again_btn", on_click=grade_card, args=(False,))
        col2.button("✅ Got it", key="fc_good_btn", on_click=grade_card, args=(True,))


def render_flashcard_tab(cfg, source):
    st.markdown("Generate a deck, browse it, drill it with spaced repetition, and export to Anki or Quizlet.")
    count = st.slider("Number of flashcards", 5, 50, 10, key="fc_count")
    if st.button("Generate flashcards", type="primary", key="fc_go"):
        if ensure_ready(source):
            task = (
                f"Create exactly {count} high-quality study flashcards covering the most important "
                "concepts, definitions, formulas, and cause-effect relationships. Each card must test one "
                "idea. Fronts are questions or prompts; backs are concise, self-contained answers. Do not "
                "repeat cards.\nReturn ONLY a JSON array with no commentary and no code fences, in this "
                'schema: [{"front": "...", "back": "...", "topic": "..."}]'
            )
            try:
                with st.spinner(f"Generating with {cfg.label}..."):
                    raw, used = call_llm(
                        cfg,
                        SYSTEM_BASE,
                        with_source(source.text, task),
                        temperature=0.4,
                        max_tokens=min(8000, 150 * count + 600),
                    )
            except LLMError as exc:
                st.error(str(exc))
            else:
                cards = [normalize_card(o) for o in find_objects(raw, {"front", "back"})]
                cards = [c for c in cards if c][:count]
                if not cards:
                    st.error("The model did not return usable flashcards. Try again or switch models.")
                    with st.expander("Raw model output"):
                        st.code(raw)
                else:
                    if len(cards) < count:
                        st.warning(f"Only {len(cards)} of {count} cards could be parsed. Output may have been truncated.")
                    st.session_state["cards"] = cards
                    st.session_state["cards_meta"] = {"model": used, "requested": cfg.model}
                    st.session_state["fc_queue"] = None
    cards = st.session_state.get("cards")
    if not cards:
        return
    meta = st.session_state["cards_meta"]
    if meta["model"] != meta["requested"]:
        st.info(f"The requested model failed, so `{meta['model']}` answered instead.")
    st.markdown(f"#### Deck ({len(cards)} cards)")
    for number, card in enumerate(cards, start=1):
        topic = f" · {card['topic']}" if card["topic"] else ""
        with st.expander(f"Card {number}{topic}: {card['front']}"):
            st.markdown(card["back"])
    render_drill(cards)
    st.markdown("#### Export")
    col1, col2 = st.columns(2)
    table = flashcards_to_table(cards)
    col1.download_button(
        "⬇️ Markdown table (.md)",
        data=build_markdown("Flashcards", table, meta["model"]),
        file_name="flashcards.md",
        mime="text/markdown",
        key="fc_dl_md",
    )
    col2.download_button(
        "⬇️ Anki/Quizlet TSV (.txt)",
        data=flashcards_to_tsv(cards),
        file_name="flashcards-anki.txt",
        mime="text/plain",
        key="fc_dl_tsv",
    )
    with st.expander("Preview Markdown table"):
        st.markdown(table, unsafe_allow_html=True)


def clear_quiz_widgets():
    for key in [k for k in st.session_state.keys() if k.startswith("quiz_q_")]:
        del st.session_state[key]


def retake_quiz():
    clear_quiz_widgets()
    st.session_state["quiz_submitted"] = False
    st.session_state["quiz_answers"] = None


def render_quiz_results(quiz):
    answers = st.session_state["quiz_answers"]
    total = len(quiz)
    chosen_indexes = []
    for i, q in enumerate(quiz):
        labels = [option_label(k, t) for k, t in enumerate(q["options"])]
        picked = answers[i]
        chosen_indexes.append(labels.index(picked) if picked in labels else None)
    score = sum(1 for i, q in enumerate(quiz) if chosen_indexes[i] == q["correct"])
    pct = score / total * 100
    if pct >= 90:
        verdict, color = "Outstanding", "#1e8e3e"
    elif pct >= 80:
        verdict, color = "Strong", "#1e8e3e"
    elif pct >= 60:
        verdict, color = "Passing: review the gaps below", "#f29900"
    else:
        verdict, color = "Needs more revision", "#d93025"
    st.markdown(
        f"<div style='border-left:8px solid {color};padding:14px 18px;border-radius:8px;"
        f"background:rgba(128,128,128,0.08);'><div style='font-size:2rem;font-weight:700;color:{color};'>"
        f"{score}/{total}</div><div>{pct:.0f}% · {verdict}</div></div>",
        unsafe_allow_html=True,
    )
    st.progress(score / total)
    unanswered = sum(1 for c in chosen_indexes if c is None)
    if unanswered:
        st.warning(f"{unanswered} question(s) were left unanswered and count as incorrect.")
    for i, q in enumerate(quiz):
        chosen = chosen_indexes[i]
        is_correct = chosen == q["correct"]
        icon = "✅" if is_correct else "❌"
        with st.expander(f"{icon} Question {i + 1}", expanded=True):
            st.markdown(f"**{q['question']}**")
            if chosen is None:
                st.markdown("_You did not answer this question._")
            for k, text in enumerate(q["options"]):
                if k == q["correct"]:
                    flag = "✅ **Correct answer**"
                elif k == chosen:
                    flag = "❌ **Your answer (incorrect)**"
                else:
                    flag = "▫️ Incorrect option"
                if k == chosen and k == q["correct"]:
                    flag = "✅ **Your answer (correct)**"
                rationale = q["rationales"][k] or "No rationale provided."
                st.markdown(f"**{option_label(k, text)}** — {flag}\n\n> {rationale}")
    st.button("Retake exam", key="quiz_retake", on_click=retake_quiz)
    meta = st.session_state["quiz_meta"]
    st.download_button(
        "⬇️ Download graded exam (.md)",
        data=build_markdown(f"Practice Exam: {score}/{total}", quiz_to_markdown(quiz, answers), meta["model"]),
        file_name="practice-exam-graded.md",
        mime="text/markdown",
        key="quiz_dl_graded",
    )


def render_quiz_tab(cfg, source):
    st.markdown("Generate a practice exam, answer it, then get instant grading with rationales for every option.")
    col1, col2 = st.columns(2)
    count = col1.slider("Number of questions", 3, 25, 5, key="quiz_count")
    difficulty = col2.selectbox("Difficulty", list(QUIZ_DIFFICULTY.keys()), key="quiz_difficulty")
    if st.button("Generate exam", type="primary", key="quiz_go"):
        if ensure_ready(source):
            task = (
                f"Create exactly {count} multiple-choice questions at {QUIZ_DIFFICULTY[difficulty]}.\n"
                "Rules:\n"
                "- Each question has exactly 4 options written as plain text with no letter prefixes.\n"
                "- Exactly one option is correct. Distractors must be plausible and based on real misconceptions.\n"
                "- Never use 'all of the above', 'none of the above', or 'both A and B' options.\n"
                "- Base every question on the source material only.\n"
                "- 'correct' is the single letter A, B, C or D.\n"
                "- 'rationales' is an array of 4 strings in the same order as 'options'. Each explains step by "
                "step why that option is correct or why it is wrong. Show the working for calculations.\n"
                "Return ONLY a JSON array with no commentary and no code fences, in this schema:\n"
                '[{"question": "...", "options": ["...", "...", "...", "..."], "correct": "B", '
                '"rationales": ["...", "...", "...", "..."]}]'
            )
            try:
                with st.spinner(f"Generating with {cfg.label}..."):
                    raw, used = call_llm(
                        cfg,
                        SYSTEM_BASE,
                        with_source(source.text, task),
                        temperature=0.4,
                        max_tokens=min(16000, 700 * count + 800),
                    )
            except LLMError as exc:
                st.error(str(exc))
            else:
                parsed = [normalize_question(o) for o in find_objects(raw, {"question", "options"})]
                quiz = [q for q in parsed if q][:count]
                if not quiz:
                    st.error("The model did not return a parseable exam. Try again or switch models.")
                    with st.expander("Raw model output"):
                        st.code(raw)
                else:
                    if len(quiz) < count:
                        st.warning(
                            f"Only {len(quiz)} of {count} questions were valid. The output may have been truncated "
                            "or some questions were malformed."
                        )
                    clear_quiz_widgets()
                    st.session_state["quiz"] = quiz
                    st.session_state["quiz_meta"] = {"model": used, "requested": cfg.model}
                    st.session_state["quiz_submitted"] = False
                    st.session_state["quiz_answers"] = None
    quiz = st.session_state.get("quiz")
    if not quiz:
        return
    meta = st.session_state["quiz_meta"]
    if meta["model"] != meta["requested"]:
        st.info(f"The requested model failed, so `{meta['model']}` answered instead.")
    if st.session_state["quiz_submitted"] and st.session_state["quiz_answers"] is not None:
        render_quiz_results(quiz)
        return
    with st.form("exam_form"):
        for i, q in enumerate(quiz):
            st.markdown(f"**Q{i + 1}. {q['question']}**")
            st.radio(
                f"Answer for question {i + 1}",
                [option_label(k, t) for k, t in enumerate(q["options"])],
                index=None,
                key=f"quiz_q_{i}",
                label_visibility="collapsed",
            )
            st.divider()
        submitted = st.form_submit_button("Submit Exam", type="primary")
    if submitted:
        st.session_state["quiz_answers"] = [st.session_state.get(f"quiz_q_{i}") for i in range(len(quiz))]
        st.session_state["quiz_submitted"] = True
        st.rerun()
    st.download_button(
        "⬇️ Download ungraded exam with answer key (.md)",
        data=build_markdown("Practice Exam", quiz_to_markdown(quiz), meta["model"]),
        file_name="practice-exam.md",
        mime="text/markdown",
        key="quiz_dl_blank",
    )


def render_notes_tab(cfg, source):
    st.markdown(
        "Turn sparse lecture slides or bulleted handouts into textbook-quality chapter notes. "
        "Works best on PPTX uploads."
    )
    depth = st.selectbox(
        "Expansion depth",
        ["Standard chapter notes", "Detailed textbook chapter (longer)"],
        key="notes_depth",
    )
    if st.button("Synthesize notes", type="primary", key="notes_go"):
        if ensure_ready(source):
            length = (
                "Write in full paragraphs, with worked examples for every major idea."
                if depth.startswith("Detailed")
                else "Be thorough but economical."
            )
            task = (
                "Convert these lecture slides or bullet points into textbook-quality study notes.\n"
                f"{length}\n"
                "Structure:\n"
                "1. Chapter title and learning objectives.\n"
                "2. Sections that follow the slide order, each with a heading. Expand every bullet into explanatory "
                "narrative that defines terms, explains reasoning, and connects ideas. Reference slide numbers like "
                "(Slide 4) where the source provides them.\n"
                "3. All formulas in LaTeX, with each symbol defined and one worked example.\n"
                "4. A glossary of key terms.\n"
                "5. Anticipated exam essay prompts (at least 4), each with a bullet outline of a strong answer.\n"
                "Mark any content that goes beyond what the slides state with [Added context] so the student can "
                "tell lecture content from supplementary explanation."
            )
            generate_text(cfg, "notes", SYSTEM_BASE, with_source(source.text, task), max_tokens=7000, temperature=0.35)
    show_text_result("notes", "Lecture Notes", "lecture-notes", "notes_dl")


def render_paper_tab(cfg, source):
    st.markdown("Analyse a research paper or journal article: methodology, findings, critique, and citations.")
    if st.button("Analyse paper", type="primary", key="paper_go"):
        if ensure_ready(source):
            task = (
                "Analyse this research paper. Use exactly these sections:\n"
                "## 1. Bibliographic Metadata\nTitle, authors, venue/journal, year, volume/issue/pages, DOI. "
                "Write 'Not found in document' for anything not present. Never guess.\n"
                "## 2. Core Research Question\n"
                "## 3. Hypothesis\nState whether it is explicit or inferred by you.\n"
                "## 4. Methodology Breakdown\nDesign, sample (size, population, sampling method), controls, "
                "independent/dependent/confounding variables, instruments or datasets, analysis methods.\n"
                "## 5. Key Findings\nInclude the actual numbers, effect sizes and p-values reported.\n"
                "## 6. Critique and Limitations\nStrengths, internal and external validity threats, statistical "
                "concerns, sampling bias, overreach in the conclusions, and what the authors do not address.\n"
                "## 7. Follow-up Questions\n3 to 5 research questions that would strengthen or test the work.\n"
                "## 8. Citations\nProvide an APA 7th edition reference and an IEEE reference built only from the "
                "metadata in section 1. Where a field is missing, insert a bracketed placeholder such as [year] "
                "and tell the student to verify against the publisher's page. Never invent a DOI."
            )
            generate_text(cfg, "paper", SYSTEM_BASE, with_source(source.text, task), max_tokens=5000, temperature=0.2)
    show_text_result("paper", "Research Paper Analysis", "paper-analysis", "paper_dl")


def render_project_tab(cfg, source):
    st.markdown("Break a term project or lab assignment into architecture, phased milestones, and deliverables.")
    objective = st.text_area(
        "Project objective",
        height=100,
        key="project_objective",
        placeholder="e.g. Build a real-time sentiment dashboard for campus social media",
    )
    scope = st.text_area(
        "Scope, constraints and grading criteria",
        height=100,
        key="project_scope",
        placeholder="e.g. Team of 3, Python only, must include evaluation on a public dataset",
    )
    col1, col2 = st.columns(2)
    timeline = col1.selectbox("Timeline", PROJECT_TIMELINES, index=2, key="project_timeline")
    team = col2.number_input("Team size", min_value=1, max_value=12, value=3, key="project_team")
    st.caption("Any uploaded or pasted material (assignment brief, rubric) is used as additional context.")
    if st.button("Build roadmap", type="primary", key="project_go"):
        if not objective.strip():
            st.warning("Describe the project objective first.")
        else:
            task = (
                f"PROJECT OBJECTIVE: {objective.strip()}\n"
                f"SCOPE AND CONSTRAINTS: {scope.strip() or 'Not specified; state your assumptions.'}\n"
                f"TIMELINE: {timeline}\nTEAM SIZE: {team}\n\n"
                "Produce a project plan with these sections:\n"
                "1. Assumptions and a one-paragraph project summary.\n"
                "2. System architecture: components, responsibilities, data flow, and an ASCII diagram in a "
                "code block.\n"
                "3. Phased roadmap scaled to the timeline, with week ranges, tasks, owner suggestions for the "
                "team size, and exit criteria for each phase: Phase 1 Specs & Architecture, Phase 2 Core "
                "Development, Phase 3 Testing & Evaluation, Phase 4 Report & Defense.\n"
                "4. Deliverables checklist using '- [ ]' items, mapped to phases.\n"
                "5. Recommended tech stack and tools as a table with the reason for each choice and one "
                "simpler fallback.\n"
                "6. Risk register: top 5 risks with mitigations.\n"
                "7. Evaluation plan and report/defense outline.\n"
                "Be realistic about what a student team can finish in the given time."
            )
            if source.text.strip():
                user = with_source(source.text, task)
            else:
                user = task
            generate_text(cfg, "project", SYSTEM_BASE, user, max_tokens=6000, temperature=0.4)
    show_text_result("project", "Term Project Roadmap", "project-roadmap", "project_dl")


def main():
    st.set_page_config(page_title=APP_TITLE, page_icon="🎓", layout="wide")
    inject_css()
    init_state()
    cfg, max_chars = render_sidebar()
    st.title("🎓 OmniStudy AI")
    st.caption("University Academic & Research Suite")
    source = render_ingestion(max_chars)
    render_metrics(cfg, source)
    st.subheader("2. Study tools")
    tabs = st.tabs(
        [
            "📝 Summary & Study Guides",
            "🃏 Flashcard Studio",
            "🧪 Practice MCQ Exam",
            "📑 Slide-to-Notes",
            "🔬 Research Paper Analyzer",
            "🛠️ Term Project Builder",
        ]
    )
    with tabs[0]:
        render_summary_tab(cfg, source)
    with tabs[1]:
        render_flashcard_tab(cfg, source)
    with tabs[2]:
        render_quiz_tab(cfg, source)
    with tabs[3]:
        render_notes_tab(cfg, source)
    with tabs[4]:
        render_paper_tab(cfg, source)
    with tabs[5]:
        render_project_tab(cfg, source)


if __name__ == "__main__":
    main()
